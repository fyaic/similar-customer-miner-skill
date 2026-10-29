"""Outcome tests for reviewable recall and scoped evidence; no network/real keys."""
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from src import qualification as q, tradewind as tw
from src.exporter import export_results
from src.models import CustomerProfile, SimilarCompany


def review(**changes):
    return {'seed_id': 'seed', 'domain': 'buyer.example', 'product': 'hand tools', 'decision': 'relevant',
            'source_url': 'https://buyer.example/products', 'excerpt': 'Hand tools catalogue for workshops',
            'checked_at': datetime.now(timezone.utc).isoformat(), 'reviewer': 'fixture reviewer',
            'reason': 'Official product category corroborates target', **changes}


def load(tmp_path, items):
    path = tmp_path/'reviews.json'
    path.write_text(json.dumps({'schema_version': 1, 'reviews': items}))
    return q.load_reviews(path, {}), path


def candidate(text='Equipment supplier'):
    return SimilarCompany(company_name='Buyer', website='https://buyer.example', main_products=text)


def seed(tags=None, id='seed'):
    return CustomerProfile(customer_id=id, product_tags=tags or ['hand tools'])


@pytest.mark.parametrize('text', ['', 'Financial tools for consumers', 'Tools to keep expenses in hand',
                                 'We do not sell hand tools', 'We no longer manufacture hand tools'])
def test_missing_or_ambiguous_text_is_reviewable_not_irrelevant(text):
    assert q.assess_product(candidate(text), seed())[0] == 'needs_review'


def test_ordered_plural_and_hyphenated_phrase_matches():
    assert q.assess_product(candidate('Professional hand-tool catalogue'), seed())[0] == 'matched'


def test_aliases_are_explicit_and_scoped():
    c = candidate('Distributor of water taps')
    assert q.assess_product(c, seed(['faucet']))[0] == 'needs_review'
    assert q.assess_product(c, seed(['faucet']), {'faucet': ['water taps']})[0] == 'matched_alias'
    assert q.assess_product(candidate('Thread-cutting taps'), seed(['faucet']), {'faucet': ['water taps']})[0] == 'needs_review'


@pytest.mark.parametrize('changes', [
    {'source_url': 'https://other.example/products'}, {'source_url': 'http://buyer.example/products'},
    {'source_url': 'https://user:pass@buyer.example/products'}, {'reviewer': ''}, {'excerpt': 'Software platform'},
    {'excerpt': 'We do not sell hand tools'}, {'decision': 'verified'}, {'checked_at': '2026-10-29T10:00:00'},
    {'checked_at': (datetime.now(timezone.utc)-timedelta(days=91)).isoformat()},
    {'checked_at': (datetime.now(timezone.utc)+timedelta(days=1)).isoformat()},
])
def test_invalid_evidence_cannot_promote(tmp_path, changes):
    with pytest.raises(ValueError): load(tmp_path, [review(**changes)])


def test_duplicates_rejected(tmp_path):
    with pytest.raises(ValueError, match='Duplicate'): load(tmp_path, [review(), review(decision='unrelated')])


def test_review_only_applies_to_same_seed_domain_and_product(tmp_path):
    reviews, _ = load(tmp_path, [review()])
    assert q.assess_product(candidate(), seed(), reviews=reviews)[0] == 'reviewed_match'
    assert q.assess_product(candidate(), seed(id='different'), reviews=reviews)[0] == 'needs_review'
    assert q.assess_product(candidate(), seed(['faucet']), reviews=reviews)[0] == 'needs_review'
    c = candidate();c.website = 'https://other.example'
    assert q.assess_product(c, seed(), reviews=reviews)[0] == 'needs_review'


def test_negative_review_is_product_specific_and_overrides_provider_text(tmp_path):
    reviews, _ = load(tmp_path, [review(decision='unrelated', reason='Category was discontinued')])
    assert q.assess_product(candidate('Hand tools'), seed(), reviews=reviews)[0] == 'reviewed_unrelated'
    assert q.assess_product(candidate('Hand tools'), seed(['hand tools', 'faucet']), reviews=reviews)[0] == 'needs_review'
    assert q.assess_product(candidate('Faucet manufacturer'), seed(['hand tools', 'faucet']), reviews=reviews)[0] == 'matched'


def test_role_claim_does_not_confirm_import_or_buying_intent():
    c = candidate('Manufacturer, distributor and importer of tools')
    q.apply_business_role(c)
    assert 'importer_claim' in c.business_role and 'distributor' in c.business_role
    assert c.buyer_qualification == 'unverified' and c.import_evidence_status == 'not_checked'


def test_negated_role_is_not_an_importer_hint():
    c = candidate('We are not an importer. We are a retailer.')
    q.apply_business_role(c)
    assert c.business_role == 'retailer'


def test_customs_same_name_with_conflicting_domain_is_not_import_proof():
    c = candidate()
    class Session:
        def call(self, *args):
            return {'operation_id': 'synthetic', 'response': {'records': [{
                'buyer': {'name': 'Buyer', 'domain': 'other.example'}, 'hs_code': '8205'}]}}
    tw.verify_customs(c, Session(), {'customs_start': '2026-01-01', 'customs_end': '2026-10-01',
                                   'customs_importer_country': 'GB', 'customs_exporter_country': 'CN'})
    assert c.hs_codes == [] and c.import_evidence_status == 'no_exact_match'
    assert c.buyer_qualification == 'unverified'


@pytest.fixture
def run_case(tmp_path, monkeypatch):
    monkeypatch.setattr(tw, 'settings', lambda *a: {'key': 'fixture', 'base': 'https://app.trade-wind.co',
        'insight_prefix': '/api/insight', 'state': str(tmp_path/'state')})
    def run(opts=None, detail=None, search=None):
        cfg = {'tradewind': {'max_credits': '5', 'max_enrich': 1, 'enrich_people': False, **(opts or {})},
               'pipeline': {'target_countries': ['GB'], 'max_results_per_customer': 3}}
        calls = []
        def send(method, url, body, headers, timeout):
            calls.append(url)
            response = {'success': True, 'request_id': 'synthetic', 'meta': {'billing': {'credits_charged': .24}}}
            company = {'name': 'Buyer', 'domain': 'buyer.example', 'country': 'GB', 'description': 'Equipment supplier'}
            if url.endswith('/company/search'): response['companies'] = search if search is not None else [company]
            elif url.endswith('/company/enrich'): response['company'] = detail or company
            elif url.endswith('/people/search'): response['people'] = []
            else: raise AssertionError(url)
            return 200, response
        rows = tw.mine([seed()], cfg, tmp_path/'result.xlsx', send)
        return rows, cfg, calls
    return run


def test_unknown_retained_in_visible_review_sheet_without_contact_spend(run_case, tmp_path):
    rows, cfg, calls = run_case()
    assert rows == [] and not any('/people/' in url for url in calls)
    assert cfg['_run_stats']['tradewind']['review_queue_count'] == 1
    assert cfg['_run_stats']['tradewind']['excluded_count'] == 0
    assert json.loads((tmp_path/'result.review.json').read_text())['candidates'][0]['website'] == 'https://buyer.example'
    export_results(rows, str(tmp_path/'result.xlsx'), cfg)
    book = load_workbook(tmp_path/'result.xlsx')
    assert book.sheetnames == ['线索名单', '待复核', 'Results', 'Summary']
    assert book['待复核'].cell(3, 2).value == 'Buyer'
    assert book['Results'].max_row == 1


def test_official_review_recovers_candidate_but_not_buyer_status(run_case, tmp_path):
    _, path = load(tmp_path, [review()])
    rows, cfg, calls = run_case({'product_evidence_file': str(path)})
    assert len(rows) == 1 and rows[0].product_match_status == 'reviewed_match'
    assert rows[0].product_evidence_citation == 'https://buyer.example/products'
    assert rows[0].product_evidence_reviewer == 'fixture reviewer'
    assert rows[0].buyer_qualification == 'unverified' and not rows[0].email_verified
    assert cfg['_run_stats']['tradewind']['review_queue_count'] == 0
    assert any('/people/search' in url for url in calls)


def test_product_proof_never_overrides_country_conflict(run_case, tmp_path):
    _, path = load(tmp_path, [review()])
    rows, cfg, calls = run_case({'product_evidence_file': str(path)},
        detail={'name': 'Buyer', 'domain': 'buyer.example', 'country': 'US'})
    assert rows == [] and cfg['_run_stats']['tradewind']['excluded_count'] == 1
    assert cfg['_run_stats']['tradewind']['review_queue_count'] == 0
    assert not any('/people/' in url for url in calls)


def test_invalid_review_rejected_before_credentials_or_network(tmp_path, monkeypatch):
    path = tmp_path/'bad.json';path.write_text('{}')
    monkeypatch.setattr(tw, 'settings', lambda *a: pytest.fail('Credentials loaded'))
    with pytest.raises(ValueError):
        tw.mine([seed()], {'tradewind': {'product_evidence_file': str(path)}}, tmp_path/'unused.xlsx')


def test_reviewed_unrelated_never_gets_contact_spend_in_exploration(run_case, tmp_path):
    _, path = load(tmp_path, [review(decision='unrelated')])
    rows, cfg, calls = run_case({'product_evidence_file': str(path), 'require_product_match': False})
    assert rows == [] and cfg['_run_stats']['tradewind']['excluded_count'] == 1
    assert len(calls) == 1


def test_qualification_only_retains_product_candidate_without_contact_calls(run_case, tmp_path):
    _, path = load(tmp_path, [review()])
    rows, _, calls = run_case({'product_evidence_file': str(path), 'contact_search_enabled': False})
    assert len(rows) == 1 and not any('/people/' in url for url in calls)


def test_negative_review_does_not_consume_company_enrichment_slot(run_case, tmp_path):
    _, path = load(tmp_path, [review(domain='excluded.example', source_url='https://excluded.example/products', decision='unrelated')])
    _, cfg, calls = run_case({'product_evidence_file': str(path)}, search=[
        {'name': 'Excluded', 'domain': 'excluded.example', 'country': 'GB', 'description': 'Hand tools manufacturer'},
        {'name': 'Buyer', 'domain': 'buyer.example', 'country': 'GB'}])
    assert sum(url.endswith('/company/enrich') for url in calls) == 1
    assert cfg['_review_queue'][0]['company_name'] == 'Buyer'
    assert cfg['_run_stats']['tradewind']['excluded_count'] == 1
