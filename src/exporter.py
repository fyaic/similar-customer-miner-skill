"""Export SimilarCompany results to Excel/CSV/Markdown."""
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd

from src import table_style as ts
from src.models import SimilarCompany
from tools.email_tool import is_generic_email

logger = logging.getLogger("similar_customer_miner")


def export_results(
    companies: List[SimilarCompany],
    output_path: str,
    config: Dict[str, Any],
) -> str:
    """Export results to the requested format.

    xlsx layout (product-level readability):
      1. 线索名单  — curated, styled, human-facing lead list (first tab)
      2. Results   — full field-level value/citation/confidence/notes (audit/validation)
      3. Summary   — counts
    """
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    rows = [_to_row(c) for c in companies]
    df = pd.DataFrame(rows)

    if path.suffix.lower() == ".xlsx":
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Results", index=False)
            # Styled, human-facing lead list as the first tab.
            _write_leads_sheet(writer.book, companies)
            if config.get("output", {}).get("include_summary_sheet", True):
                _write_summary_sheet(writer.book, companies, config)
            _style_results_sheet(writer.book)
            # Whole-table font constraint: force every used cell to 微软雅黑.
            ts.enforce_font(writer.book)
        logger.info("Exported %d companies to %s", len(companies), output_path)
        return str(path)
    elif path.suffix.lower() == ".csv":
        df.to_csv(path, index=False, encoding="utf-8-sig")
    elif path.suffix.lower() == ".md":
        path.write_text(df.to_markdown(index=False), encoding="utf-8")
    else:
        raise ValueError(f"Unsupported output format: {path.suffix}")

    logger.info("Exported %d companies to %s", len(companies), output_path)
    return str(path)


def _to_row(company: SimilarCompany) -> Dict[str, Any]:
    row: Dict[str, Any] = {
        "source_seed": company.source_seed,
        "priority_score": company.priority_score,
        "fit_score": company.fit_score,
        "similarity_score": company.similarity_score,
        "overall_confidence": company.overall_confidence,
        "grade": company.grade,
        "fact_check_status": company.fact_check_status,
        "datasource": (company.all_citations or {}).get("datasource"),
    }

    fields = [
        "company_name", "country", "website", "main_products", "hs_codes", "industry_tags",
        "revenue", "customs_china_purchase", "estimated_shipping_spend", "shipment_frequency",
        "contact_name", "contact_title",
        "phone", "email",
        "linkedin", "facebook", "twitter", "youtube",
    ]

    is_demo = bool(company.notes and "demo fallback" in company.notes.lower())

    for field in fields:
        value = getattr(company, field)
        if isinstance(value, list):
            value = "; ".join(str(v) for v in value)
        row[f"{field}_value"] = value
        row[f"{field}_citation"] = getattr(company, f"{field}_citation", None)
        row[f"{field}_confidence"] = getattr(company, f"{field}_confidence", 0.0)
        row[f"{field}_notes"] = _field_notes(company, field, is_demo)

    row["procurement_proxy_notes_value"] = company.procurement_proxy_notes
    row["email_verified"] = company.email_verified
    row["notes"] = company.notes
    row["all_citations"] = str(company.all_citations) if company.all_citations else None
    return row


def _field_notes(company: SimilarCompany, field: str, is_demo: bool) -> Optional[str]:
    """Return a short, human-readable note for fields that are inferred or unverified."""
    confidence = getattr(company, f"{field}_confidence", 0.0) or 0.0
    value = getattr(company, field)
    if value is None or (isinstance(value, list) and not value):
        return None

    labels: List[str] = []
    if is_demo:
        labels.append("demo fallback")

    if field in ("contact_name", "contact_title") and confidence < 0.5:
        labels.append("inferred (not a verified decision maker)")

    if field == "email" and confidence < 0.5:
        labels.append("inferred from domain pattern")

    if field in ("linkedin", "facebook", "twitter", "youtube") and confidence < 0.6:
        if company.notes and "static public pointers" in company.notes:
            labels.append("unverified static pointer")
        else:
            labels.append("unverified")

    return "; ".join(labels) if labels else None


def _write_summary_sheet(workbook, companies: List[SimilarCompany], config: Optional[Dict[str, Any]] = None) -> None:
    """Styled Summary with counts + a mandatory A/B/C grade explanation block."""
    config = config or {}
    thr = config.get("pipeline", {}).get("grade_thresholds", {})
    ga = thr.get("A", 0.80)
    gb = thr.get("B", 0.50)
    counts = {"A": 0, "B": 0, "C": 0}
    for c in companies:
        counts[c.grade] = counts.get(c.grade, 0) + 1
    passed = sum(1 for c in companies if c.fact_check_status == "passed")

    ws = workbook.create_sheet("Summary")
    ncols = 4

    # Block 1: 统计概览
    ts.section_title(ws, 1, ncols, "统计概览")
    ts.header_row(ws, 2, ["指标", "数值", "", ""])
    stats = [
        ("候选公司总数", len(companies)),
        ("Grade A", counts.get("A", 0)),
        ("Grade B", counts.get("B", 0)),
        ("Grade C", counts.get("C", 0)),
        ("通过事实核查", passed),
    ]
    r = 3
    for label, val in stats:
        ts.data_cell(ws, r, 1, label)
        ts.data_cell(ws, r, 2, val, align="center")
        ts.data_cell(ws, r, 3, "")
        ts.data_cell(ws, r, 4, "")
        r += 1

    # Block 2: 评级(Grade)说明 —— 每次输出必须包含
    r += 1  # blank separator
    ts.section_title(ws, r, ncols, "评级 (Grade) 说明")
    r += 1
    ts.header_row(ws, r, ["评级", "判定阈值 (综合置信度)", "含义", "建议动作"])
    r += 1
    grade_rows = [
        ("A", f"≥ {ga:.2f}", "关键字段齐全且可信（来源明确、置信度高）", "可直接触达"),
        ("B", f"{gb:.2f} – {ga:.2f}", "主要字段可信，部分信息需人工复核", "复核后触达"),
        ("C", f"< {gb:.2f}", "信息稀疏或待验证（缺联系人/发货量等）", "仅作参考，人工补全后再触达"),
    ]
    grade_bg = {"A": ts.GRADE_A_BG, "B": ts.GRADE_B_BG, "C": ts.GRADE_C_BG}
    for g, thr_txt, meaning, action in grade_rows:
        ts.data_cell(ws, r, 1, g, bg=grade_bg[g], align="center", bold=True)
        ts.data_cell(ws, r, 2, thr_txt, align="center")
        ts.data_cell(ws, r, 3, meaning)
        ts.data_cell(ws, r, 4, action)
        r += 1

    # Note: grade != priority
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    note = ws.cell(row=r, column=1, value=(
        "说明：评级(Grade) 只表示单条数据的「可信度」，不代表触达「优先级」。"
        "触达先后请看「线索名单」第一列 priority_score（综合相似度 + 进口强度 + 可触达完整度）。"
    ))
    note.font = ts.Font(color=ts.SECONDARY_FONT, size=9, italic=True)
    note.alignment = ts.Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = 30

    ts.set_widths(ws, [10, 20, 44, 24])
    ts.freeze_header(ws, first_data_row=3)


# --- Results (full audit) sheet: header style + semantic column widths --------

def _results_col_width(name: str) -> int:
    """Wide for name/website/products/citations/notes; narrow for scores."""
    n = str(name)
    if n.endswith("_confidence"):
        return 11
    if n.endswith("_citation"):
        return 34   # citations are usually URLs — keep readable
    if n.endswith("_notes"):
        return 24
    if n in ("notes", "all_citations"):
        return 44
    wide_value = {
        "website_value": 38,
        "company_name_value": 28,
        "main_products_value": 34,
        "email_value": 30,
        "estimated_shipping_spend_value": 30,
        "procurement_proxy_notes_value": 34,
        "linkedin_value": 30, "facebook_value": 30, "twitter_value": 26, "youtube_value": 26,
        "contact_name_value": 18, "contact_title_value": 20,
        "hs_codes_value": 16, "industry_tags_value": 18,
    }
    if n in wide_value:
        return wide_value[n]
    if n.endswith("_value"):
        return 18
    if n in ("source_seed", "datasource", "fact_check_status"):
        return 16
    if n in ("priority_score", "fit_score", "similarity_score", "overall_confidence", "grade", "email_verified"):
        return 13
    return 16


def _style_results_sheet(workbook) -> None:
    """Style the full-detail Results sheet: green header, semantic widths, wrap,
    and Chinese user-friendly header labels."""
    if "Results" not in workbook.sheetnames:
        return
    ws = workbook["Results"]
    from openpyxl.utils import get_column_letter
    from src.field_labels import label_for
    # header row: set width by canonical name, then relabel to Chinese for display
    for cell in ws[1]:
        canonical = cell.value
        ws.column_dimensions[get_column_letter(cell.column)].width = _results_col_width(canonical)
        cell.value = label_for(canonical)
        cell.fill = ts.fill(ts.HEADER_BG)
        cell.font = ts.Font(color=ts.HEADER_FONT, bold=True, size=10)
        cell.alignment = ts.Alignment(horizontal="center", vertical="center", wrap_text=True)
    # data cells: wrap + top align so wide key columns show full content
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = ts.Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


# --- Styled, human-facing lead list ------------------------------------------

_LEADS_HEADERS = [
    "优先级", "评级", "匹配度", "公司名", "官网", "HS / 主营",
    "联系人", "邮箱", "电话", "社媒", "发货量 / 采购(proxy)", "数据来源", "备注",
]
_LEADS_WIDTHS = [8, 7, 8, 26, 30, 16, 18, 30, 15, 22, 24, 12, 34]


def _write_leads_sheet(workbook, companies: List[SimilarCompany]) -> None:
    """Build a curated, product-level lead list as the first tab."""
    ws = workbook.create_sheet("线索名单", 0)
    ncols = len(_LEADS_HEADERS)

    ts.section_title(ws, 1, ncols, f"相似买家线索名单 · 共 {len(companies)} 家 · 按优先级排序")
    ts.header_row(ws, 2, _LEADS_HEADERS)

    r = 3
    for c in companies:
        grade_bg = {"A": ts.GRADE_A_BG, "B": ts.GRADE_B_BG}.get(c.grade, ts.GRADE_C_BG)
        email = c.email or ""
        email_secondary = bool(email) and (is_generic_email(email) or (c.email_confidence or 0) < 0.5)
        social = " ".join(u for u in (c.linkedin, c.facebook, c.twitter, c.youtube) if u)
        contact = c.contact_name or ""
        if c.contact_name and c.contact_title:
            contact = f"{c.contact_name}\n{c.contact_title}"
        hs_or_prod = "; ".join(c.hs_codes) if c.hs_codes else (c.main_products or "")

        # procurement column: shipment frequency + proxy shipping-spend (warning tint)
        proc_parts = [p for p in (c.shipment_frequency, c.estimated_shipping_spend) if p]
        proc_text = "\n".join(proc_parts)
        proc_bg = ts.WARNING_BG if c.estimated_shipping_spend else None

        ts.data_cell(ws, r, 1, _pct(c.priority_score), align="center", bold=True)
        ts.data_cell(ws, r, 2, c.grade, bg=grade_bg, align="center", bold=True)
        ts.data_cell(ws, r, 3, _pct(c.fit_score), align="center")
        ts.data_cell(ws, r, 4, c.company_name, bold=True)
        _website_cell(ws, r, 5, c.website)
        ts.data_cell(ws, r, 6, hs_or_prod)
        ts.data_cell(ws, r, 7, contact or "—", secondary=not contact)
        ts.data_cell(ws, r, 8, email or "—", secondary=(email_secondary or not email))
        ts.data_cell(ws, r, 9, c.phone or "—", secondary=not c.phone)
        ts.data_cell(ws, r, 10, social or "—", secondary=not social)
        ts.data_cell(ws, r, 11, proc_text or "—", bg=proc_bg, secondary=not proc_text)
        ts.data_cell(ws, r, 12, (c.all_citations or {}).get("datasource", ""), align="center", secondary=True)
        ts.data_cell(ws, r, 13, c.notes or "", secondary=True)
        r += 1

    ts.set_widths(ws, _LEADS_WIDTHS)
    ts.freeze_header(ws, first_data_row=3)


def _website_cell(ws, row: int, col: int, website: Optional[str]) -> None:
    ts.data_cell(ws, row, col, website or "—", secondary=not website)
    if website:
        cell = ws.cell(row=row, column=col)
        cell.hyperlink = website
        cell.font = ts.Font(color="1155CC", size=10, underline="single")


def _pct(score: Optional[float]) -> Optional[int]:
    return round((score or 0) * 100)
