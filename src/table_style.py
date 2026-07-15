"""Product-level table styling for xlsx output.

Design language ported from the `tommy-table-canvas` business-mapping style
(green section titles, light-green headers, alternating rows, semantic tints,
wrap-everything, frozen header). Adapted from Tencent Docs payloads to openpyxl
so the local .xlsx reads like a finished product, not a raw dump.
"""
from typing import List, Optional

from openpyxl.styles import Alignment, Border, Font as _Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Whole-table font constraint: every cell must render in Microsoft YaHei. ---
FONT_NAME = "微软雅黑"


def Font(*args, **kwargs) -> _Font:
    """openpyxl Font wrapper that forces the mandated table font family.

    All styling in this module (and the exporter) builds fonts through this, so
    the whole used range is guaranteed to be 微软雅黑 regardless of call site.
    """
    kwargs.setdefault("name", FONT_NAME)
    return _Font(*args, **kwargs)


# --- Palette (hex, no leading FF; openpyxl fills/fonts take RRGGBB) ---
SECTION_TITLE_BG = "319B62"   # green
SECTION_TITLE_FONT = "FFFFFF"
HEADER_BG = "EAFAF1"          # light green
HEADER_FONT = "1F1F1F"
DATA_FONT = "1F1F1F"
SECONDARY_FONT = "8C8C8C"     # citations / low-confidence notes
ROW_ODD_BG = "FFFFFF"
ROW_EVEN_BG = "FAFAFA"
WARNING_BG = "FFF7E6"         # proxy / unverified
SUCCESS_BG = "F6FFED"         # grade A / verified
ERROR_BG = "FFF1F0"
GRADE_A_BG = "D9F7BE"
GRADE_B_BG = "F6FFED"
GRADE_C_BG = "FAFAFA"
BORDER_COLOR = "E8E8E8"

_THIN = Side(style="thin", color=BORDER_COLOR)
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def fill(hexcolor: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hexcolor)


def section_title(ws, row: int, ncols: int, text: str, height: int = 30) -> None:
    """Merged, centered green banner across the block width."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill = fill(SECTION_TITLE_BG)
    cell.font = Font(color=SECTION_TITLE_FONT, bold=True, size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = fill(SECTION_TITLE_BG)
    ws.row_dimensions[row].height = height


def header_row(ws, row: int, headers: List[str], height: int = 28) -> None:
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=text)
        cell.fill = fill(HEADER_BG)
        cell.font = Font(color=HEADER_FONT, bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def data_cell(
    ws,
    row: int,
    col: int,
    value,
    *,
    zebra: bool = True,
    bg: Optional[str] = None,
    secondary: bool = False,
    bold: bool = False,
    align: str = "left",
) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    if bg is None and zebra:
        bg = ROW_EVEN_BG if row % 2 == 0 else ROW_ODD_BG
    if bg:
        cell.fill = fill(bg)
    cell.font = Font(
        color=SECONDARY_FONT if secondary else DATA_FONT,
        size=10,
        bold=bold,
        italic=secondary,
    )
    cell.alignment = Alignment(horizontal=align, vertical="top", wrap_text=True)
    cell.border = BORDER


def set_widths(ws, widths: List[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def freeze_header(ws, first_data_row: int) -> None:
    """Freeze the title + header rows so they stay visible when scrolling."""
    ws.freeze_panes = f"A{first_data_row}"


def enforce_font(workbook, font_name: str = FONT_NAME) -> None:
    """Hard constraint: force EVERY used cell in EVERY sheet to `font_name`.

    Run this last, after pandas has written its sheets (which default to Calibri),
    so the whole workbook is guaranteed 微软雅黑 while preserving each cell's other
    font attributes (size / bold / italic / color / underline).
    """
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                f = cell.font
                cell.font = _Font(
                    name=font_name,
                    size=f.size,
                    bold=f.bold,
                    italic=f.italic,
                    color=f.color,
                    underline=f.underline,
                    strike=f.strike,
                    vertAlign=f.vertAlign,
                )
